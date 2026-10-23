from pathlib import Path
import json,re,hashlib
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
R=Path(__file__).resolve().parent;out=R/'deliverables';out.mkdir(exist_ok=True)
files={'pl':['01_grudzien_bez_pospiechu_PL.md','02_kalendarz_z_zadaniami_PL.md'],'en':['01_december_family_time_EN.md','02_advent_activities_EN.md']}
receipts=[]
def runs(p,t):
 for j,x in enumerate(re.split(r'(\*\*.*?\*\*)',t)):
  if not x:continue
  run=p.add_run(x[2:-2] if x.startswith('**') else x);run.bold=x.startswith('**')
for lang,names in files.items():
 d=Document();sec=d.sections[0];sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
 sec.top_margin=Inches(.72);sec.bottom_margin=Inches(.7);sec.left_margin=sec.right_margin=Inches(.78)
 for name,size,bold,color in [('Normal',10.5,False,'202124'),('Title',19,True,'000000'),('Heading 1',13,True,'422B50'),('Heading 2',10.5,True,'422B50')]:
  st=d.styles[name];st.font.name='DejaVu Sans';st.font.size=Pt(size);st.font.bold=bold;st.font.color.rgb=RGBColor.from_string(color)
  st.paragraph_format.line_spacing=1.04;st.paragraph_format.space_after=Pt(4 if name=='Normal' else 6)
  if name!='Normal':st.paragraph_format.space_before=Pt(8 if name!='Title' else 0);st.paragraph_format.keep_with_next=True
 normal=d.styles['Normal'];normal.paragraph_format.widow_control=True
 for index,f in enumerate(names):
  if index:d.add_page_break()
  source=(R/'candidate'/f).read_text();body=re.sub(r'\A---\n.*?\n---\n','',source,flags=re.S).strip()
  for block in body.split('\n\n'):
   block=block.strip()
   if not block:continue
   if block.startswith('# '):p=d.add_paragraph(block[2:],style='Title')
   elif block.startswith('## '):p=d.add_paragraph(block[3:],style='Heading 1')
   elif block.startswith('### '):p=d.add_paragraph(block[4:],style='Heading 2')
   else:p=d.add_paragraph();runs(p,block)
   p.paragraph_format.widow_control=True
 foot=sec.footer.paragraphs[0];foot.alignment=2;f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');foot._p.append(f)
 d.core_properties.author='Rise.Shine.Evolve.';d.core_properties.title='Święta są też po drodze Artykuły' if lang=='pl' else 'Before Christmas Slips By Articles'
 d.core_properties.subject='Two complete articles';d.core_properties.comments=''
 for tree in [d._element,d.styles.element]:
  for border in list(tree.iter(qn('w:pBdr'))):border.getparent().remove(border)
 path=out/('SWIETA_SA_TEZ_PO_DRODZE_ARTYKULY_PL_V4.docx' if lang=='pl' else 'BEFORE_CHRISTMAS_SLIPS_BY_ARTICLES_EN_V4.docx');d.save(path)
 # Each reader paragraph must round-trip the exact candidate prose, no editorial metadata.
 expected=[]
 for f in names:
  body=re.sub(r'\A---\n.*?\n---\n','',(R/'candidate'/f).read_text(),flags=re.S).strip()
  expected.extend(re.sub(r'^#{1,3} ','',p).replace('**','') for p in body.split('\n\n') if p.strip())
 actual=[p.text for p in Document(path).paragraphs if p.text.strip()];assert actual==expected
 receipts.append({'language':lang,'docx':str(path),'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'candidate_sha256':hashlib.sha256((R/'candidate'/('candidate-'+lang+'.json')).read_bytes()).hexdigest(),'exact_paragraph_roundtrip':'PASS','paragraphs':len(expected),'render_visual':'PENDING'})
(R/'qa/DOCX_BINDING.json').write_text(json.dumps(receipts,ensure_ascii=False,indent=2)+'\n');print(json.dumps(receipts))
